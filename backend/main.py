"""FastAPI backend của MedAI Diagnostic Portal.

Mỗi route trong file này đại diện cho một chức năng nghiệp vụ chính:
- quản lý bệnh nhân
- upload ảnh X-quang
- gọi mô hình AI dự đoán
- lưu phiên chụp vào SQLite
- xuất báo cáo Word cho người dùng

File này đóng vai trò "bộ não" của ứng dụng, còn database.py là tầng lưu trữ
và ai_model.py là tầng xử lý tri thức / mô hình.
"""

import io
import uuid
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Literal

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field

from . import database
from .ai_model import CheckpointInference, ModelUnavailable, localized_labels


BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
DATA_DIR = ROOT_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
MAX_IMAGE_BYTES = 10 * 1024 * 1024
model: CheckpointInference | None = None
model_error = ""


@asynccontextmanager
async def lifespan(_: FastAPI):
    """Khởi tạo ứng dụng khi server start.

    Chức năng chính:
    - đảm bảo thư mục uploads tồn tại
    - khởi tạo SQLite nếu chưa có bảng
    - cố gắng nạp mô hình AI nếu checkpoint tồn tại
    """
    global model, model_error
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    database.init_db()
    try:
        model = CheckpointInference()
    except (ModelUnavailable, ImportError) as error:
        model = None
        model_error = str(error)
    yield


app = FastAPI(title="MedAI Diagnostic Portal Demo", lifespan=lifespan)


def present_exam(exam: dict) -> dict:
    """Định dạng dữ liệu phiên chụp để frontend dễ render.

    Tại đây, labels gốc là dictionary số điểm theo tên bệnh, còn
    labels_localized là phiên bản đã chuyển sang tiếng Việt và mức độ cảnh báo.
    """
    return {**exam, "labels_localized": localized_labels(exam["labels"])}


class PatientCreate(BaseModel):
    """Schema dữ liệu để tạo bệnh nhân mới."""

    medical_record_number: str = Field(min_length=2, max_length=30)
    full_name: str = Field(min_length=2, max_length=120)
    date_of_birth: str
    gender: Literal["Nữ", "Nam", "Khác"]
    phone: str = Field(default="", max_length=30)


class ExamUpdate(BaseModel):
    """Schema dữ liệu để cập nhật findings, impression và trạng thái duyệt."""

    findings: str = Field(max_length=5000)
    impression: str = Field(max_length=3000)
    approved: bool = False


@app.get("/", include_in_schema=False)
def home() -> FileResponse:
    """Serve giao diện web chính cho người dùng."""
    return FileResponse(ROOT_DIR / "frontend" / "index.html")


@app.get("/api/patients")
def patients() -> list[dict]:
    """Trả về danh sách bệnh nhân cho sidebar và dashboard."""
    return database.list_patients()


@app.post("/api/patients", status_code=201)
def add_patient(payload: PatientCreate) -> dict:
    """Tạo hồ sơ bệnh nhân mới."""
    try:
        return database.create_patient(**payload.model_dump())
    except Exception as error:
        if "UNIQUE constraint failed" in str(error):
            raise HTTPException(status_code=409, detail="Mã hồ sơ đã tồn tại") from error
        raise


@app.get("/api/patients/{patient_id}/exams")
def patient_exams(patient_id: int) -> list[dict]:
    """Lấy tất cả phiên chụp của một bệnh nhân."""
    if not database.get_patient(patient_id):
        raise HTTPException(status_code=404, detail="Không tìm thấy bệnh nhân")
    return [present_exam(exam) for exam in database.list_exams(patient_id)]


@app.get("/api/exams/{exam_id}/images/{view}")
def exam_image(exam_id: int, view: Literal["frontal", "lateral"]) -> FileResponse:
    """Trả về ảnh frontal hoặc lateral đã lưu trong thư mục uploads."""
    exam = database.get_exam(exam_id)
    if not exam:
        raise HTTPException(status_code=404, detail="Không tìm thấy phiên chụp")
    filename = exam[f"{view}_path"]
    image_path = (UPLOAD_DIR / filename).resolve()
    if image_path.parent != UPLOAD_DIR.resolve() or not image_path.is_file():
        raise HTTPException(status_code=404, detail="Không tìm thấy ảnh")
    return FileResponse(image_path)


async def read_image(upload: UploadFile) -> bytes:
    """Xác thực và đọc nội dung file ảnh từ request upload.

    Hành vi chính:
    - kiểm tra MIME type
    - giới hạn dung lượng ảnh tối đa 10 MB
    - trả về bytes để lưu xuống đĩa
    """
    allowed_types = {"image/png", "image/jpeg", "image/webp"}
    if upload.content_type not in allowed_types:
        raise HTTPException(status_code=415, detail="Chỉ nhận ảnh PNG, JPG hoặc WebP")
    content = await upload.read(MAX_IMAGE_BYTES + 1)
    if not content or len(content) > MAX_IMAGE_BYTES:
        raise HTTPException(status_code=413, detail="Ảnh phải nhỏ hơn 10 MB")
    return content


@app.post("/api/exams/analyze", status_code=201)
async def analyze_exam(
    patient_id: int = Form(...),
    frontal: UploadFile = File(...),
    lateral: UploadFile = File(...),
) -> dict:
    """Phân tích 2 ảnh X-quang của bệnh nhân và lưu kết quả vào database.

    Quy trình:
    1. kiểm tra bệnh nhân tồn tại
    2. đọc ảnh frontal/lateral và validate
    3. lưu ảnh xuống data/uploads
    4. gọi model.predict() để lấy điểm độ nguy cơ
    5. gọi model.draft_report() để sinh findings/impression
    6. lưu exam vào SQLite và trả JSON cho frontend
    """
    if not database.get_patient(patient_id):
        raise HTTPException(status_code=404, detail="Không tìm thấy bệnh nhân")
    frontal_bytes = await read_image(frontal)
    lateral_bytes = await read_image(lateral)
    if model is None:
        raise HTTPException(status_code=503, detail=f"AI chưa sẵn sàng: {model_error}")
    extensions = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp"}
    frontal_name = f"{uuid.uuid4().hex}_frontal{extensions[frontal.content_type]}"
    lateral_name = f"{uuid.uuid4().hex}_lateral{extensions[lateral.content_type]}"
    (UPLOAD_DIR / frontal_name).write_bytes(frontal_bytes)
    (UPLOAD_DIR / lateral_name).write_bytes(lateral_bytes)
    labels = model.predict(frontal_bytes, lateral_bytes)
    findings, impression = model.draft_report(labels, frontal_bytes, lateral_bytes)
    exam = database.create_exam(
        patient_id,
        frontal_name,
        lateral_name,
        labels,
        findings,
        impression,
    )
    return present_exam(exam)


@app.put("/api/exams/{exam_id}")
def edit_exam(exam_id: int, payload: ExamUpdate) -> dict:
    """Cập nhật findings/impression và trạng thái phê duyệt của phiên chụp."""
    exam = database.update_exam(
        exam_id, payload.findings, payload.impression, payload.approved
    )
    if not exam:
        raise HTTPException(status_code=404, detail="Không tìm thấy phiên chụp")
    return present_exam(exam)


@app.get("/api/exams/{exam_id}/report.docx")
def download_report(exam_id: int) -> StreamingResponse:
    """Xuất báo cáo Word cho một phiên chụp đã có."""
    exam = database.get_exam(exam_id)
    if not exam:
        raise HTTPException(status_code=404, detail="Không tìm thấy phiên chụp")
    patient = database.get_patient(exam["patient_id"])
    document = Document()
    document.add_heading("MEDAI | BÁO CÁO X-QUANG NGỰC", 0)
    document.add_paragraph("BẢN DEMO - KHÔNG DÙNG ĐỂ CHẨN ĐOÁN HOẶC ĐIỀU TRỊ")
    document.add_heading("Thông tin bệnh nhân", level=1)
    document.add_paragraph(f"Họ tên: {patient['full_name']}")
    document.add_paragraph(f"Mã hồ sơ: {patient['medical_record_number']}")
    document.add_paragraph(f"Ngày sinh: {patient['date_of_birth']}")
    document.add_paragraph(f"Thời gian: {exam['created_at']}")
    document.add_heading("Findings", level=1)
    document.add_paragraph(exam["findings"])
    document.add_heading("Impression", level=1)
    document.add_paragraph(exam["impression"])
    document.add_heading("Điểm dự đoán của mô hình (chưa kiểm định lâm sàng)", level=1)
    for item in localized_labels(exam["labels"]):
        document.add_paragraph(
            f"{item['name_vi']} - mức điểm tham khảo: {item['level']} "
            f"({item['score']:.1%})",
            style="List Bullet",
        )
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="medai-report-{exam_id}.docx"'},
    )
